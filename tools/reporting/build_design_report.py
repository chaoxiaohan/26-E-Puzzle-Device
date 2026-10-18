"""Build the editable report from measured evidence and reviewed source."""
import json
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def main():
    import argparse
    parser=argparse.ArgumentParser()
    parser.add_argument('--task', type=Path, required=True)
    args=parser.parse_args()
    task=args.task
    assets=Path(__file__).resolve().parents[2]/'docs/benchmarks/2026-10-18'
    data=json.loads((task/'summary.json').read_text())
    rows=data['cases']
    document=Document()
    section=document.sections[0]
    section.page_width=Inches(8.5);section.page_height=Inches(11)
    section.top_margin=Inches(.64);section.bottom_margin=Inches(.64)
    section.left_margin=Inches(.72);section.right_margin=Inches(.72)
    section.header_distance=Inches(.28);section.footer_distance=Inches(.28)
    for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption','Header','Footer'):
        style=document.styles[name]
        style.font.name='Calibri';style.font.color.rgb=RGBColor(0,0,0)
        fonts=style.element.get_or_add_rPr().rFonts
        fonts.attrib.clear()
        for key in ('ascii','hAnsi','cs'):fonts.set(qn('w:'+key),'Calibri')
        fonts.set(qn('w:eastAsia'),'微软雅黑' if name!='Normal' else '宋体')
        style.font.size=Pt(11)
        style.font.italic=False;style.font.bold=False
    normal=document.styles['Normal'].paragraph_format
    normal.space_after=Pt(7);normal.line_spacing=1.28
    for name,size in (('Title',30),('Subtitle',15),('Heading 1',20),('Heading 2',13),('Caption',9)):
        document.styles[name].font.size=Pt(size)
    for name in ('Heading 1','Heading 2'):
        document.styles[name].font.bold=True
        document.styles[name].paragraph_format.space_before=Pt(12)
        document.styles[name].paragraph_format.space_after=Pt(7)
    header=section.header.paragraphs[0]
    header.text='拼图装置视觉系统设计报告     Raspberry Pi 5'
    header.runs[0].font.size=Pt(8)
    footer=section.footer.paragraphs[0]
    footer.text='2026 年 10 月 18 日    图片回放与当前算法版本'
    footer.paragraph_format.tab_stops.add_tab_stop(Inches(6.55))
    footer.add_run('\t')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    for run in footer.runs:run.font.size=Pt(8)
    def p(text='',style=None):return document.add_paragraph(text,style)
    def heading(text):document.add_heading(text,2)
    def page(title):
        document.add_heading(title,1).paragraph_format.page_break_before=True
    def picture(path,width=6.95,caption=None):
        para=p();para.paragraph_format.space_after=Pt(4)
        para.alignment=WD_ALIGN_PARAGRAPH.CENTER
        run=para.add_run();shape=run.add_picture(str(path),width=Inches(width))
        shape._inline.docPr.set('descr',caption or path.stem)
        if caption:
            para.paragraph_format.keep_with_next=True
            cap=p(caption,'Caption');cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
    def table(headers,values,widths=None,size=10):
        t=document.add_table(rows=1,cols=len(headers));t.autofit=False
        if widths is None:widths=[7.06/len(headers)]*len(headers)
        for column,width in zip(t.columns,widths):column.width=Inches(width)
        border=OxmlElement('w:tblBorders')
        for side in ('top','left','bottom','right','insideH','insideV'):
            element=OxmlElement('w:'+side);element.set(qn('w:val'),'single');element.set(qn('w:sz'),'4');element.set(qn('w:color'),'D9D9D9');border.append(element)
        t._tbl.tblPr.append(border)
        for i,text in enumerate(headers):t.rows[0].cells[i].text=str(text)
        repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for values_row in values:
            cells=t.add_row().cells
            for i,text in enumerate(values_row):cells[i].text=str(text)
        for ri,row in enumerate(t.rows):
            no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
            for ci,cell in enumerate(row.cells):
                cell.width=Inches(widths[ci])
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'24465B' if ri==0 else 'F3F6F8' if ri%2==0 else 'FFFFFF');cell._tc.get_or_add_tcPr().append(shade)
                margins=OxmlElement('w:tcMar')
                for side in ('top','bottom','left','right'):
                    x=OxmlElement('w:'+side);x.set(qn('w:w'),'70');x.set(qn('w:type'),'dxa');margins.append(x)
                cell._tc.get_or_add_tcPr().append(margins)
                for para in cell.paragraphs:
                    para.paragraph_format.space_after=Pt(0);para.paragraph_format.line_spacing=1.12
                    for run in para.runs:
                        run.font.name='Calibri';run.font.size=Pt(size);run.font.bold=ri==0
                        run.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '172A3A')
                        run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
        p().paragraph_format.space_after=Pt(1)
        return t
    p('拼图装置视觉系统\n设计报告','Title')
    p('树莓派 5 算法实现与绿色实拍图测试','Subtitle')
    p('2026 年电赛 E 题工程版本    2026 年 10 月 18 日')
    heading('摘要')
    p('本系统面向少量平面碎片的识别、矩形重组与执行坐标输出，采用绿色背景分割、多帧轮廓稳定、几何搜索及扑克图案联合评分。针对相同矩形无法仅靠外形区分的问题，候选搜索保留每片身份与朝向，通过拼缝、角标、主体及源图断口证据进行排序，并保留歧义判断。')
    p('在树莓派 5 上对 43 张绿色背景实拍图完成三轮主测试，并对 10 张无印刷图补充扑克负样本测试，共 139 次有效运行。33 张扑克的总分析时间中位数为 0.973 秒，单次最大值为 6.717 秒；31 张历史扑克在三轮中的主预览与既有参考完全一致。18 张扑克通过自动执行规则，14 张存在歧义，1 张仅提供几何候选。普通模式仍有漏检与保底布局问题，全文保留这些结果。')
    picture(assets/'pairs/G31.png',caption='图 1  四片相似矩形的扑克重建示例 G31  原图与算法输出均来自实测')
    p('关键词  绿色背景分割  平面拼图  图案连续性  角标结构  树莓派 5')
    p('报告范围：当前视觉与上位机软件设计、板端图片回放及接口说明。本文不将离线预览作为机械抓取放置完成的证据。',style='Caption')

    page('阅读导航与验证范围')
    table(['章节','内容','页码'],[
        ('1','任务要求与系统方案','3'),('2','采集分割与稳定识别','4'),('3','模板与布局候选搜索','5'),
        ('4','图案评分和自动执行门槛','6'),('5','运行调度与电气接口','7'),('6','测试方法与可复现性','8'),
        ('7','测试结果与样本分析','9'),('8','逐图运行时间分布','10'),('9','逐张数据表','11 至 12'),
        ('10','问题分析与后续验证','13'),('附录 A','全部原图与拼接图','14 至 35')],[.65,5.25,1.16])
    heading('本报告能支持的结论')
    p('报告说明当前代码如何处理图片，给出同一块树莓派、同一算法和固定历史图像上的运行耗时、输出状态与图像一致性。每张图都有原始 PNG、输出 PNG、三次运行记录和来源校验值。')
    heading('需要另外验证的结论')
    p('实时视频的曝光、稳定等待和端到端交互延迟不在本次计时内。未知牌组、复杂遮挡、全光照条件、毫米级放置精度、整机两分钟完成率及机械连续运行可靠性尚无本次实测证据。')
    table(['证据类型','本次状态'],[
        ('原始照片重新检测与求解','已在真实树莓派 5 完成'),('历史扑克主预览与状态一致性','31 张 × 3 次全部一致'),
        ('实际相机与串口只读查询','沿用 2026 年 10 月 16 日记录'),('实际抓取旋转放置','本次未执行'),
        ('精确逐片物理真值','未提供，不计算物理位姿正确率')],[2.3,4.76])

    page('1 任务要求与系统方案')
    p('项目附带的 E 题赛题要求将 A4 区域中的碎片移动到目标区域并形成矩形；扑克任务还要求图案对应正确。文档给出的整机限时包括识别、规划与机构动作，不能用单次求解时间代替整机成绩。[1]')
    table(['方案','优势','适用限制'],[
        ('已学习模板匹配','目标已知时直接寻找实体对应','需要学习目标布局'),('纯几何拼接','无需印刷图案，适合普通碎片','相同外形不能唯一决定牌面'),
        ('几何与图案联合搜索','同时比较身份、朝向与邻接','依赖可见墨迹，保留歧义判断')],[1.55,2.65,2.86])
    p('系统将三种方法保留为不同运行模式。FIRST 使用已学习模板；NORMAL 以矩形覆盖和短边保护处理普通碎片；CARD 以几何约束筛选候选并联合评价印刷结构。当前实现没有神经网络或 OCR 推理过程。[2]')
    picture(assets/'figures/pipeline.png',width=6.9,caption='图 2  从图像采集到执行接口的数据流及本次计时边界')
    p('硬件职责划分为 USB 相机采集、树莓派图像处理与 Flutter 交互、USB 串口传输，以及下位机执行机构控制。报告不推断未提供的电机驱动、电源或吸盘电路参数。')

    page('2 采集分割与稳定识别')
    heading('2 1 图像与 ROI')
    p('hardware_devices.py 优先枚举 UVC 采集节点，过滤元数据和内部 ISP/编码节点。主程序处理图像方向、ROI、断线重连及照片保存。原始照片保留完整像素，另存 ROI 和检测设置；回放按每张照片自己的参数复现，不覆盖现场设置。')
    heading('2 2 自适应绿色分割')
    p('detect_pieces 保留原始 RGB ROI 用于绿色识别，另对亮度分支执行 Gamma 提亮、CLAHE、Gaussian 模糊以及 Otsu/手动阈值。自动绿色检测以三像素步长采样，在 HSV 与绿色通道占比约束下定位主颜色簇，由色相中值、偏差和饱和度/亮度分位数确定范围。')
    p('绿色背景与纸片掩膜经过按配置启用的形态学处理。在线输入颜色状态时，稳定颜色簇采用旧值 0.82、新值 0.18 的平滑；单张回放没有跨帧状态。采样不足时退回固定绿色规则。')
    heading('2 3 从轮廓到实体描述')
    p('外轮廓通过面积、ROI 边缘距离、尺寸和实心度过滤。复杂连通轮廓会尝试分离；多边形简化与长直边拟合降低锯齿对角点的影响。每片输出顶点、中心、面积、实心度和长边朝向，最多保留四片。编号代表实体身份，不能直接等同于左上或右下的位置。')
    p('扑克路径针对四片场景的四边形进一步拟合原始直边，检查角点漂移、面积比例和顶点方向后才接受修正。该步骤属于测量改进，不把不同实体合并为同一种形状。')
    heading('2 4 视频中的轮廓稳定')
    p('主循环保留最近两秒检测，从最近八帧中选取相同顶点数量组合，枚举实体对应，限制中心位移不超过 12 像素。最多使用五组、至少三组，对齐顶点后取中值，结合 MAD 方式剔除离群样本，并检查角点离散度及面积波动。')
    p('本次只测试独立静态图片，不复制同一照片来伪造多帧稳定效果。因此照片回放与在线取中值轮廓可能存在细小差异；2026 年 10 月 16 日的实时结果与本次两张照片状态不同，不能据此认定算法回退。')

    page('3 模板与布局候选搜索')
    heading('3 1 FIRST 已知目标模式')
    p('模板学习保存图像块、掩膜、轮廓、中心和面积。匹配组合 Hu 形状相似度、轮廓姿态对应、面积比和纹理相关度，权重为 0.30、0.40、0.20、0.10；四片时穷举最多 24 种一一分配。匹配后采用刚体变换回到学习位置。本次没有逐例学习数据，不计算模板模式命中率。')
    heading('3 2 NORMAL 几何模式')
    p('原生 C++ 将多边形缩放到栅格，用扫描线构造 64 位位集，快速计算覆盖、重叠和候选组合。11 个宽高比由两个工作线程分配搜索，外层按填充率和位置步长逐级尝试。NORMAL 与界面一致开启短边保护。')
    p('候选还需经过连续坐标的覆盖、接边、边界直线和重叠检查。严格矩形不成立时，程序可能返回 BEST 等无重叠保底布局，或继续 Python/其他几何构造；这些状态不等价于满足全部赛题要求。')
    heading('3 3 CARD 联合模式')
    table(['阶段','当前处理'],[
        ('输入检查','验证 2 至 4 片、多边形有效性和绿色结构适用条件'),
        ('优先候选','由真实纸角构建最多 24 个候选；结构弱时补充扇形和条带切法'),
        ('其他拓扑','无候选时按矩形分类尝试联合拓扑或限时边邻接图搜索'),
        ('几何兜底','无可用布局或印刷结构弱时调用 C++ 种子，并重排实体身份与方向'),
        ('统一比较','几何预筛、图案评分、等价整体姿态去重、首选与竞争方案比较')],[1.2,5.86])
    p('相同矩形的边长可能完全一致，但每片携带的墨迹不同。当前搜索保留单片编号和顶点对应，只消除整幅布局共同的旋转和平移；单片半转仍是不同候选。候选池上限、分数上界与时间预算共同限制计算量。该过程不是任意切法的全局最优证明。')

    page('4 图案评分和自动执行门槛')
    heading('4 1 局部拼缝与整牌结构')
    p('拼缝沿边采样 96 个位置，向内取六层像素。比较器主要利用最近两层，提取灰度、暗度、饱和度、红色、黑色和梯度，一次比较七个沿边错位量。绿色边缘被剔除，无墨迹白边保持低证据，避免“都为白色”被误判为正确拼接。')
    table(['证据','实现','作用'],[
        ('局部拼缝','绝对连续性与双向相对匹配','识别邻接与朝向'),('外边界和断口','追踪源墨迹组件的切口采样','防止半个图案跑到外侧'),
        ('角标','字形与花色组件的方向和布局','区分相似碎片的实体位置'),('主体与文字串','重复图案、排列及细长组件关系','支持点数牌、人物牌与 JOKER')],[1.15,3.2,2.71])
    p('每片按白色参考进行颜色归一化。结构评分使用旋转和平移，拒绝反射；角标与字母串使用轮廓启发式，不执行 OCR。基础分以纹理为主，结构置信度足够时再与结构分按 0.45/0.55 融合。分数用于排序，不代表概率。')
    heading('4 2 从首选到自动执行')
    p('基础准入要求纹理置信度不低于 0.30、纹理分不低于 0.62、覆盖不低于 0.84、内部匹配不低于 0.55，并检查外侧参与比例与竞争差值。结构分支进一步检查墨迹、外边截断、角标错位和断口复合；强结构证据可通过另一组更严格的几何门槛。完整条件见 main.py 中 select_card_candidate。')
    table(['输出状态','含义'],[
        ('VERIFIED','通过当前自动执行规则，仍需标定和机构条件有效'),('AMBIGUOUS','有可用预览，但证据或竞争差异不足'),
        ('GEOMETRY_ONLY','几何候选或几何等级不足，不宣称牌面已确认'),('WEAK_TEXTURE / NO_LAYOUT','印刷证据不足或无布局')],[2.25,4.81])

    page('5 运行调度与电气接口')
    heading('5 1 控制计算量与界面负载')
    p('常规视觉循环上限为 10 FPS，OpenCV 显式限制为一个线程，求解由单独进程执行。求解期间停止相机/检测/JPEG 更新，保留状态心跳、界面控制和串口服务，结束后恢复画面。C++ 求解器内部使用两个线程，避免无控制地争用四个核心。')
    p('帧内复用纸片特征、掩膜、连通分量、形状相似度和对称边对缓存；结构评分可根据上界提前剪枝。候选不足时才使用几何种子。速度提升来自减少无效搜索和重复计算，CPU 使用率本身不能证明已达算法极限。')
    heading('5 2 标定和运动坐标')
    p('六组像素坐标与平台坐标用于二维仿射最小二乘拟合。程序检查点云分散程度、矩阵退化与比例，并记录 RMS 和最大残差。动作计划由源片中心、目标位置和旋转量组成。标定有效标志不等价于毫米级精度，仍需测量实际放置误差。')
    heading('5 3 硬件接口与软件边界')
    table(['接口','已知实现','本次验证边界'],[
        ('相机到树莓派','USB/UVC 图像采集','本次使用保存照片，未测曝光链路'),('树莓派到平台','CH340 USB 串口，现场波特率 115200','未发送动作命令'),
        ('通信帧','AA55 同步头、命令、标志、序号、长度、负载和 CRC16','代码和历史协议测试支持'),('下位机到机构','由原平台控制程序执行','电机/吸盘电路及误差未在本次核验')],[1.55,3.1,2.41])
    p('ACK/NACK、超时与重试保护通信过程，不能替代机械到位反馈。实际电源容量、驱动额定值、机构负载和吸附可靠性应在整机电路与机械报告中另行提供，本文不填入未经核实的参数。')
    heading('5 4 部署与回滚')
    p('树莓派使用独立发布目录与虚拟环境，旧启动路径由符号链接保留。切换前备份原程序和五份参数，保留 captures；Git 分支区分树莓派和鲁班猫。本次只增加测试工具、证据与文档，未替换板端运行程序，也未覆盖用户的 ROI 或标定配置。[4]')

    page('6 测试方法与可复现性')
    table(['项目','设置'],[
        ('设备','Raspberry Pi 5 Model B Rev 1.1，4 GB 内存，四核 ARM64'),('系统与依赖','Debian 13；Python 3.13.5；OpenCV 4.10.0；NumPy 2.2.4'),
        ('代码基线','运行代码 02505cf；算法基线 6de9e25'),('数据范围','61 份拍摄记录；选入 43 张绿底，排除 14 张黑底与 4 张遮挡图'),
        ('样本组成','33 张扑克（含 4 张 JOKER），10 张无印刷碎片（含深色片）'),('有效运行','43 张 × 3 次主测试，加 10 次 CARD 负样本，共 139 次'),
        ('执行方法','独立子进程、重新检测、原生 C++；NORMAL 开启短边保护'),('现场干扰','测试开始时 UI 与实时后端未运行；不发送串口命令')],[1.5,5.56])
    heading('6 1 计时定义')
    p('detect 为 detect_pieces 调用时间；solve 为实际 worker 调用时间，CARD 已包含候选和主预览的渲染。NORMAL 的主图在 worker 后生成，记入 post_preview。analysis 从检测前到主预览可用，包含检测、求解及这一后处理。图像解码、进程启动和文件导出分别记录或排除，不能与 analysis 混用。')
    p('每图三次独立运行，逐图表采用中位数，范围取最小至最大；全组中位数是逐图中位数的中位数。图集选择总分析时间居中的一次输出。每次保存 JSON、原图 SHA256、状态、节点数、评分、源轮廓、目标布局、CPU 时间和进程峰值 RSS。')
    heading('6 2 复核与修订')
    p('正式运行前有一张扑克预热，不纳入统计。初版普通模式回放漏开短边保护，原始 30 次记录单独归档并作废；按界面调用方式重跑的 30 次替换普通模式结果，最终有效统计仍为 139 次。测试工具修正前后均有 Git 提交。')
    p('31 张历史扑克与 10 月 16 日参考逐次核对 PNG 哈希和状态，共 93 项一致性检查。重复运行不是新增独立样本，也不是准确率的分母扩充。所有历史参数仅用于离线回放，现场设置保持原样。[3]')

    page('7 测试结果与样本分析')
    values=[]
    for key,name in (('historical_cards','历史扑克 31 张'),('cards','全部扑克 33 张'),('plain','普通碎片 10 张')):
        g=data['groups'][key]['timings_ms']
        values.append((name,f"{g['detect']['median_of_image_medians']:.1f}",f"{g['solve']['median_of_image_medians']/1000:.3f}",f"{g['analysis']['median_of_image_medians']/1000:.3f}",f"{g['analysis']['max_single_run']/1000:.3f}"))
    table(['分组','检测 ms','求解 s','总分析 s','单次最大 s'],values,[1.9,1.0,1.2,1.35,1.61])
    p('前三项时间为各图三次中位数再取组中位数。全部扑克中 18 张 VERIFIED、14 张 AMBIGUOUS、1 张 GEOMETRY_ONLY；三轮状态及主图均稳定。18/33 是当前样本的自动准入比例，不是整体识别准确率。')
    p('10 张无印刷样本在 NORMAL 中均返回 BEST，5 张漏检到不足四片，不能报告为严格矩形通过。它们在 CARD 中分别返回 3 张 WEAK_TEXTURE、2 张 GEOMETRY_ONLY 和 5 张 NO_LAYOUT，自动放行 0/10，符合不把无印刷片当作已确认扑克的预期。')
    picture(assets/'pairs/G42.png',width=6.8,caption='图 3  新增梅花 10 样本 G42  可以显示整牌，但自动判断仍为 AMBIGUOUS')
    p('两张新增梅花 10 总分析中位数分别为 6.513 秒和 6.620 秒，是本组长尾。31 张历史扑克的三轮主图均与参考逐像素一致；这一结论只证明回归稳定。观测温度 39.15 至 50.15°C，降频标记均为 0x0；单个测试进程最大 RSS 为 126.3 MiB，不是整机内存占用。')

    page('8 逐图运行时间分布')
    picture(assets/'figures/timing.png',width=6.95,caption='图 4  总分析时间的中位数及三次极差  编号与后续表格及附录一致')
    p('颜色表示输出状态：绿色为 VERIFIED，黄色为 AMBIGUOUS，红色为 GEOMETRY_ONLY，灰蓝色为普通模式 BEST。普通碎片的短耗时中包含漏检场景，不应作为完整四片求解速度优势。',style='Caption')

    for start,title in ((0,'9 逐张数据表 上'),(22,'9 逐张数据表 下')):
        page(title)
        values=[]
        for row in rows[start:start+22]:
            t=row['timings_ms']
            values.append((row['id'],row['capture'],row['piece_count'],f"{t['detect']['median']:.1f}",f"{t['solve']['median']/1000:.3f}",f"{t['analysis']['median']/1000:.3f}",row['status']))
        table(['编号','照片编号','片数','检测 ms','求解 s','总分析 s','状态'],values,[.42,1.68,.4,.7,.7,.7,2.46],size=9)
        p('表中片数为重新检测的片数，不是人工标注。时间均为三次中位数；逐次时间、极差、资源数据和来源校验值见 CSV/JSON。所有原图及输出图保留在附录 A 与可筛选 HTML 图集。',style='Caption')
        if start:
            heading('负样本与异常样本')
            p('G02、G03、G06、G07 只检出三片，G25 只检出两片。普通模式 BEST 保留的是算法当前返回结果，不能因为显示了白色矩形底板就认定完整复原。G28 为扑克 GEOMETRY_ONLY，其余待确认扑克使用 AMBIGUOUS 状态。')

    page('10 问题分析与后续验证')
    heading('10 1 当前证据所揭示的问题')
    p('普通碎片的五张漏检优先影响完整性，应检查面积过滤、小片与细长片、ROI 边缘和多边形近似，建立逐片人工标注后再优化。其余普通图虽检出四片，也返回 BEST，需继续检查严格几何约束和保底路径，不能只依据耗时评价。')
    p('扑克的部分白缝、外边界不齐和浅墨迹使候选间证据不足。新版可以处理重复矩形，但自动执行判断仍受竞争解、结构阈值和图像质量影响。建议先固定采集距离、曝光和背景，再增加未用于调参的牌组、摆放及光照，分别评估首选正确率与自动准入精确率。')
    heading('10 2 后续整机验证')
    p('逐片标注预期身份与方向，建立独立测试集；为失败样本保存完整候选。随后验证六点标定残差、不同位置的实际落点误差、旋转方向、吸附与放置可靠性。最终以连续多次完整流程计时，从解除相机遮挡至完成指示，统计整机成功率与时间分布。')
    p('本次结果说明当前树莓派算法在既有扑克回归集上输出稳定，大多数样本的软件分析时间低于数秒；同时明确暴露普通碎片完整性和部分扑克歧义问题。报告没有据此给出整机赛题得分或一般化准确率。')
    heading('参考资料与数据来源')
    p('[1] 项目随附 E题_拼图装置.pdf，2026 年全国大学生电子设计竞赛赛区赛暨模拟电子系统设计专题赛选拔赛题，第 1 至 3 页。需求取自该本地文件。',style='Caption')
    p('[2] chaoxiaohan/26-E-Puzzle-Device，raspberry-pi5 分支。运行代码基线 02505cf，算法基线 6de9e25；main.py、card_structure.py、card_face_solver.py、rectangle_solver.cpp、template_matcher.py、serial_protocol.py。完整说明：docs/VISION_ALGORITHM.md。',style='Caption')
    p('[3] 2026-10-18 板端图片回放：docs/benchmarks/2026-10-18/summary.json、per-image.csv、all-trials.csv、green-cases.json 与逐图 records；完整运行证据另存 evidence-final.tar.gz。',style='Caption')
    p('[4] docs/pi5-board-validation-20261016.json：树莓派原生构建、参数备份、真实相机、只读串口及暂停恢复记录。本次未重新验证机械动作。',style='Caption')

    for start in range(0,len(rows),2):
        page('附录 A 原始图像与拼接结果 '+str(start//2+1).zfill(2))
        p('每图依次为完整原图、对应 ROI、算法重建。像素内容未修图；预览不等同于实物拼装照片。',style='Caption')
        for row in rows[start:start+2]:
            picture(assets/'pairs'/(row['id']+'.png'),width=6.95)
            t=row['timings_ms']['analysis']
            p(f"{row['id']}  三次总分析范围 {t['min']/1000:.3f} 至 {t['max']/1000:.3f} s；展示第 {row['representative_run']} 次；进程峰值 RSS {row['peak_rss_mib']:.1f} MiB。",style='Caption')
            p('方法 '+str(row['method'])+'；三轮状态与主图一致。'+('检测不足四片，不能视为完整求解。' if row['piece_count']!=4 else ''),style='Caption')
        if start==42:
            heading('图像与数据的取用方式')
            p('HTML 图集可按模式、状态和拍摄编号筛选，并直接打开原始 PNG、拼接 PNG 与该次 JSON。CSV 适合二次统计和绘图；原始图片与输出均带 SHA256，便于核对引用是否来自同一版本。')
            p('报告插图采用等比例缩放和统一留白，不调整牌面亮度、颜色或断口，也不删除失败样本。原图可能包含黑色补边或深色碎片；筛选依据是实际工作背景为绿色，而不是碎片颜色。')
    output=task/'拼图装置视觉系统设计报告.docx'
    document.core_properties.title='拼图装置视觉系统设计报告'
    document.core_properties.subject='树莓派 5 算法实现与绿色实拍图测试'
    document.core_properties.author='拼图装置项目'
    document.core_properties.keywords='拼图;视觉;Raspberry Pi 5;绿色背景;测试'
    for root in (document.styles.element,document.element):
        for border in list(root.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    document.save(output)
    print(output)


if __name__ == '__main__':
    main()
